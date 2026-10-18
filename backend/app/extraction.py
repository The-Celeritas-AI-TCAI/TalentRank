"""Resume text extraction with per-file failure isolation and protected-attribute redaction."""
from __future__ import annotations
import re
from dataclasses import dataclass, field
from pathlib import Path

ALLOWED_EXT = {".pdf", ".docx", ".txt"}
MAX_BYTES = 10 * 1024 * 1024
MAGIC = {".pdf": b"%PDF", ".docx": b"PK"}


@dataclass
class ExtractionResult:
    filename: str
    ok: bool
    pages: list[str] = field(default_factory=list)  # one entry per page (docx/txt = 1)
    error: str | None = None

    @property
    def text(self) -> str:
        return "\n".join(self.pages)


def extract(path: str | Path) -> ExtractionResult:
    """Never raises: a bad file yields ok=False with a reason (never a zero score)."""
    p = Path(path)
    ext = p.suffix.lower()
    try:
        if ext not in ALLOWED_EXT:
            return ExtractionResult(p.name, False, error=f"Unsupported file type '{ext}'")
        size = p.stat().st_size
        if size == 0:
            return ExtractionResult(p.name, False, error="File is empty")
        if size > MAX_BYTES:
            return ExtractionResult(p.name, False, error="File exceeds 10 MB limit")
        if ext in MAGIC and not p.read_bytes()[:4].startswith(MAGIC[ext]):
            return ExtractionResult(p.name, False, error="File content does not match its extension")
        if ext == ".pdf":
            import pymupdf
            with pymupdf.open(p) as doc:
                if doc.needs_pass:
                    return ExtractionResult(p.name, False, error="PDF is password protected")
                pages = [pg.get_text() for pg in doc]
        elif ext == ".docx":
            import docx
            d = docx.Document(str(p))
            parts = [para.text for para in d.paragraphs]
            for t in d.tables:
                for row in t.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
            pages = ["\n".join(parts)]
        else:
            pages = [p.read_text(encoding="utf-8", errors="replace")]
        if len("".join(pages).strip()) < 50:
            return ExtractionResult(p.name, False, pages, "Too little extractable text (scanned image? OCR not enabled)")
        return ExtractionResult(p.name, True, pages)
    except Exception as e:  # corrupt files etc.
        return ExtractionResult(p.name, False, error=f"Could not parse file: {type(e).__name__}")


# --- Protected characteristics and proxies: removed BEFORE any scoring or AI call ---
_PROTECTED_LINE = re.compile(
    r"\b(date of birth|dob|age|gender|sex|male|female|marital status|married|single|religion|"
    r"caste|nationality|ethnicity|race|disability|disabled|pregnan\w*|photo(graph)?|"
    r"father'?s name|mother'?s name|pronouns?)\b\s*[:\-]", re.I)
_YEARS_OLD = re.compile(r"\b\d{1,2}\s*(?:years?|yrs?)\s*old\b", re.I)
_PROTECTED_PREFIX = re.compile(
    r"^\s*(?:date of birth|dob|age|gender|sex|male|female|marital status|married|single|religion|"
    r"caste|nationality|ethnicity|race|disability|disabled|pregnan\w*|photo(graph)?|"
    r"father'?s name|mother'?s name|pronouns?)\b", re.I)


def redact_protected(text: str) -> str:
    out = []
    for line in text.splitlines():
        if _PROTECTED_LINE.search(line) or _PROTECTED_PREFIX.search(line) or _YEARS_OLD.search(line):
            out.append("[REDACTED]")
        else:
            out.append(line)
    return "\n".join(out)
