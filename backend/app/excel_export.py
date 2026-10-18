"""Reference-style, filterable Excel workbook for a single resume-screening run."""
from __future__ import annotations
from datetime import datetime, timezone
from io import BytesIO
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT = "Aptos"
DANGEROUS = ("=", "+", "-", "@", "\t", "\r", "\n")
HEADERS = [
    ("Rank", 8), ("Candidate", 25), ("Resume File", 30), ("Email", 29), ("Phone", 20),
    ("Location", 25), ("Current / Latest Role", 38), ("Total Experience (yrs)", 15),
    ("Role-Specific Experience (yrs)", 18), ("Mandatory Requirement Check", 48),
    ("Selected Job Role", 24), ("Key Matched Skills", 48), ("Gaps vs Must-Haves", 48),
    ("Role-Specific Experience (20)", 17), ("Overall Relevant Experience (10)", 19),
    ("Generic Must-Haves (25)", 17), ("Role-Specific Must-Haves (30)", 20),
    ("Nice-to-Have (10)", 14), ("Evidence Quality (5)", 15), ("Total Score (100)", 15),
    ("Fit Band", 16), ("Recommendation", 34), ("Why This Score (Rationale)", 64),
    ("Red Flags / Points to Verify", 56),
]
# Backward-compatible name used by the original core test suite.
RANK_COLUMNS = HEADERS

def safe_text(value):
    if isinstance(value, str) and value and (value.startswith(DANGEROUS) or value.lstrip().startswith(DANGEROUS[:4])):
        return "'" + value
    return value

def _profile(c):
    return (c.get("criterion_evidence") or {}).get("_profile") or (c.get("criterion_gaps") or {}).get("_profile") or {}

def _mandatory_text(profile):
    items=profile.get("mandatory_check", [])
    if not items: return "No explicit mandatory requirement identified; review the job description."
    return "\n".join(f'{x.get("requirement","")}: {x.get("status","Needs review")}' for x in items)

def _recommendation(c, profile):
    if c.get("status") == "failed": return "Re-upload or manually review extraction failure"
    if c.get("status") == "incomplete": return "Manual review — insufficient evidence"
    missing=profile.get("mandatory_missing", [])
    if missing: return "Human review — verify mandatory requirements"
    score=c.get("match_score")
    if score is None: return "Manual review — score unavailable"
    if score>=80: return "Prioritize for human technical review"
    if score>=65: return "Review evidence and remaining gaps"
    if score>=50: return "Consider only after reviewing must-have gaps"
    return "Manual review required; do not auto-reject"

def build_workbook(job: dict, candidates: list[dict], ranks: dict) -> bytes:
    wb=Workbook()
    ws=wb.active; ws.title="Candidate Screening"
    navy="1F3A5F"; blue="D9EAF7"; line="D7DEE8"; thin=Side(style="thin",color=line)
    border=Border(left=thin,right=thin,top=thin,bottom=thin)
    now=datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=len(HEADERS))
    ws.cell(1,1,"TalentRank AI — Candidate Screening").font=Font(name=FONT,size=16,bold=True,color="FFFFFF")
    ws.cell(1,1).fill=PatternFill("solid",fgColor=navy)
    ws.cell(1,1).alignment=Alignment(vertical="center")
    ws.row_dimensions[1].height=30
    ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=len(HEADERS))
    ws.cell(2,1,f"Selected role: {job.get('title','')} | Exported: {now} | Scores are decision support only; human review is required.").font=Font(name=FONT,italic=True,color="596579")
    ws.cell(2,1).alignment=Alignment(wrap_text=True,vertical="top")
    ws.merge_cells(start_row=3,start_column=1,end_row=3,end_column=len(HEADERS))
    ws.cell(3,1,"Evidence rule: skills evidenced in work/project text are distinguished from skills mentioned only in a skills list. Missing information is unverified, not proof a candidate lacks the skill.").font=Font(name=FONT,size=10,color="596579")
    ws.cell(3,1).alignment=Alignment(wrap_text=True,vertical="top")
    ws.row_dimensions[2].height=26; ws.row_dimensions[3].height=30
    hdr=4
    for i,(label,width) in enumerate(HEADERS,1):
        cell=ws.cell(hdr,i,label); cell.font=Font(name=FONT,bold=True,color="FFFFFF")
        cell.fill=PatternFill("solid",fgColor=navy); cell.alignment=Alignment(wrap_text=True,vertical="center",horizontal="center"); cell.border=border
        ws.column_dimensions[get_column_letter(i)].width=width
    ws.row_dimensions[hdr].height=42
    ordered=sorted(candidates,key=lambda c:(c.get("status")!="scored", ranks.get(c.get("id")) or 10**9, -(c.get("match_score") or 0), (c.get("name") or "").lower()))
    for row_idx,c in enumerate(ordered,hdr+1):
        profile=_profile(c); breakdown=profile.get("score_breakdown",{})
        demonstrated=profile.get("demonstrated_skills",c.get("verified_skills",[]))
        listed=profile.get("listed_skills",c.get("unverified_skills",[]))
        missing=profile.get("missing_skills",c.get("confirmed_absent_skills",[]))
        matched=[]
        if demonstrated: matched.append("Demonstrated in work/projects: "+", ".join(demonstrated))
        if listed: matched.append("Listed only — verify: "+", ".join(listed))
        gaps=[]
        if missing: gaps.append("Not found in extracted text: "+", ".join(missing))
        if listed: gaps.append("Practical evidence not located for: "+", ".join(listed))
        gaps.extend(profile.get("general_gaps",[]))
        if c.get("error"): gaps.append(c["error"])
        mandatory=_mandatory_text(profile)
        score=c.get("match_score") if c.get("status")=="scored" else None
        values=[
            ranks.get(c.get("id")) if c.get("status")=="scored" else None,
            c.get("name") or "Candidate", c.get("filename",""), profile.get("email",""), profile.get("phone",""),
            profile.get("location",""), profile.get("latest_role",""),
            profile.get("total_experience"), profile.get("role_experience"), mandatory,
            job.get("title",""), "\n".join(matched) or "No JD skill evidence extracted",
            "\n".join(dict.fromkeys(gaps)) or "No explicit gaps detected by text rules; verify manually",
            breakdown.get("Role-Specific Experience (20)",0), breakdown.get("Overall Relevant Experience (10)",0),
            breakdown.get("Generic Must-Haves (25)",0), breakdown.get("Role-Specific Must-Haves (30)",0),
            breakdown.get("Nice-to-Have (10)",0), breakdown.get("Evidence Quality (5)",0),
            score, c.get("fit_band") if score is not None else "Not scored",
            _recommendation(c,profile), profile.get("rationale",c.get("evidence_summary","")),
            "\n".join(profile.get("red_flags",[])) or ("Extraction failed: "+c.get("error","") if c.get("status")=="failed" else "No automated red flag; manual verification still required"),
        ]
        for col,value in enumerate(values,1):
            cell=ws.cell(row_idx,col,safe_text(value)); cell.font=Font(name=FONT,size=10); cell.border=border
            cell.alignment=Alignment(wrap_text=True,vertical="top",horizontal="center" if col in (1,8,9,14,15,16,17,18,19,20,21) else "left")
            if col in (8,9): cell.number_format="0.0"
            if 14<=col<=20: cell.number_format="0.0"
            if c.get("status")!="scored": cell.fill=PatternFill("solid",fgColor="FFF4CE")
        ws.row_dimensions[row_idx].height=75
    last=max(hdr+1,hdr+len(ordered))
    ws.freeze_panes="C5"
    ws.auto_filter.ref=f"A4:{get_column_letter(len(HEADERS))}{last}"
    score_col=20
    for lo,hi,color in [(80,100,"C6EFCE"),(65,79.999,"DDEBF7"),(50,64.999,"FFEB9C"),(0,49.999,"F8CBAD")]:
        ws.conditional_formatting.add(f"{get_column_letter(score_col)}5:{get_column_letter(score_col)}{last}",CellIsRule(operator="between",formula=[str(lo),str(hi)],fill=PatternFill("solid",fgColor=color)))
    for col in range(14,20):
        ws.conditional_formatting.add(f"{get_column_letter(col)}5:{get_column_letter(col)}{last}",DataBarRule(start_type="num",start_value=0,end_type="num",end_value=col in (14,15,16,17) and [20,10,25,30][col-14] or (10 if col==18 else 5),color="5B9BD5"))
    # Summary worksheet
    summary=wb.create_sheet("Summary")
    summary.merge_cells("A1:D1"); summary["A1"]="TalentRank AI — Screening Summary"
    summary["A1"].font=Font(name=FONT,size=16,bold=True,color="FFFFFF"); summary["A1"].fill=PatternFill("solid",fgColor=navy); summary.row_dimensions[1].height=30
    summary["A3"]="Selected Job Role"; summary["B3"]=safe_text(job.get("title",""))
    summary["A4"]="Resumes submitted"; summary["B4"]=len(candidates)
    summary["A5"]="Successfully scored"; summary["B5"]=sum(c.get("status")=="scored" for c in candidates)
    summary["A6"]="Extraction failures"; summary["B6"]=sum(c.get("status")=="failed" for c in candidates)
    summary["A7"]="Incomplete assessments"; summary["B7"]=sum(c.get("status")=="incomplete" for c in candidates)
    for r in range(3,8): summary.cell(r,1).font=Font(name=FONT,bold=True); summary.cell(r,2).font=Font(name=FONT)
    for col,val in enumerate(["Fit Band","Score Range","Candidate Count","Suggested HR Action"],1):
        c=summary.cell(10,col,val); c.font=Font(name=FONT,bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor=navy); c.alignment=Alignment(wrap_text=True,vertical="center")
    band_actions=[("Strong Fit","80–100","Prioritize for human technical review"),("Good Fit","65–79","Review evidence and verify remaining gaps"),("Moderate Fit","50–64","Hold for human review against must-haves"),("Not a Fit","Below 50","Review evidence manually; do not auto-reject")]
    for r,(band,rng,action) in enumerate(band_actions,11):
        count=sum(c.get("status")=="scored" and c.get("fit_band")==band for c in candidates)
        for col,val in enumerate([band,rng,count,action],1):
            cell=summary.cell(r,col,val); cell.font=Font(name=FONT); cell.alignment=Alignment(wrap_text=True,vertical="top"); cell.border=border
    summary["A17"]="Mandatory requirement review"; summary["A17"].font=Font(name=FONT,bold=True,color=navy)
    summary["A18"]="Candidates with mandatory items unverified or missing"
    summary["B18"]=sum(bool(_profile(c).get("mandatory_missing")) for c in candidates if c.get("status")=="scored")
    summary["A20"]="Decision-support note"; summary["A20"].font=Font(name=FONT,bold=True,color=navy)
    summary.merge_cells("A21:D23"); summary["A21"]="Scores are based on deterministic text matching and explicit resume evidence, not a validated predictor of job performance. Do not make automated rejection or hiring decisions. Review the source resume and verify all mandatory requirements, dates, qualifications and claims with the candidate."
    summary["A21"].alignment=Alignment(wrap_text=True,vertical="top"); summary["A21"].font=Font(name=FONT,color="596579")
    for col,width in {"A":42,"B":24,"C":18,"D":54}.items(): summary.column_dimensions[col].width=width
    summary.freeze_panes="A10"
    # Rubric worksheet
    rubric=wb.create_sheet("Rubric")
    rubric.merge_cells("A1:C1"); rubric["A1"]="Scoring Rubric & Job Description Reference"
    rubric["A1"].font=Font(name=FONT,size=15,bold=True,color="FFFFFF"); rubric["A1"].fill=PatternFill("solid",fgColor=navy)
    rubric["A3"]="Component"; rubric["B3"]="Maximum Points"; rubric["C3"]="How points are assigned"
    for c in rubric[3]:
        c.font=Font(name=FONT,bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor=navy); c.alignment=Alignment(wrap_text=True)
    rubric_rows=[
      ("Role-specific experience",20,"Only an explicit years statement tied to the selected role/domain is scored. No duration is inferred from dates."),
      ("Overall relevant experience",10,"Uses only an explicit total-years-of-experience statement. Compared with an explicit years threshold in the JD, or a five-year reference when no threshold is supplied."),
      ("Generic must-haves",25,"Skills named in mandatory JD language and identified as general transferable skills. Work/project evidence earns full credit; skills-list-only mention earns 40% credit; not found earns 0."),
      ("Role-specific must-haves",30,"Mandatory technical/domain skills from the JD. Work/project evidence earns full credit; skills-list-only mention earns 40% credit; not found earns 0."),
      ("Nice-to-have",10,"Preferred/desirable JD skills, using the same evidence distinction."),
      ("Evidence quality",5,"Small fixed rubric for presence of clearly labelled work history, skills section, explicit experience statement and identifiable latest role."),
    ]
    for r,row in enumerate(rubric_rows,4):
        for col,val in enumerate(row,1):
            c=rubric.cell(r,col,safe_text(val)); c.font=Font(name=FONT); c.alignment=Alignment(wrap_text=True,vertical="top"); c.border=border
    rubric["A12"]="Fit Band"; rubric["B12"]="Score"; rubric["C12"]="Suggested interpretation"
    for c in rubric[12]: c.font=Font(name=FONT,bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor=navy)
    bands=[("Strong Fit","80–100","Review first; verify all requirements"),("Good Fit","65–79","Review matched evidence and gaps"),("Moderate Fit","50–64","Manual review against must-haves"),("Not a Fit","Below 50","Low text match; never auto-reject")]
    for r,row in enumerate(bands,13):
        for col,val in enumerate(row,1):
            c=rubric.cell(r,col,val); c.font=Font(name=FONT); c.alignment=Alignment(wrap_text=True,vertical="top"); c.border=border
    rubric["A19"]="Selected role"; rubric["B19"]=safe_text(job.get("title",""))
    rubric["A20"]="Job description / summary"; rubric["B20"]=safe_text(job.get("description",""))
    rubric.merge_cells("B20:C24"); rubric["B20"].alignment=Alignment(wrap_text=True,vertical="top")
    rubric["A26"]="Fairness & limitations"; rubric["B26"]="Protected attributes (age, gender, religion, caste, disability, marital status, ethnicity and similar characteristics) are redacted before scoring. Scores are deterministic decision support only. Extraction failures are reported separately; no score is fabricated."
    rubric.merge_cells("B26:C28"); rubric["B26"].alignment=Alignment(wrap_text=True,vertical="top")
    rubric.column_dimensions["A"].width=32; rubric.column_dimensions["B"].width=25; rubric.column_dimensions["C"].width=90
    rubric.freeze_panes="A4"
    for sheet in wb.worksheets:
        sheet.sheet_view.showGridLines=False
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is not None and cell.alignment == Alignment():
                    cell.alignment=Alignment(vertical="top",wrap_text=True)
    out=BytesIO(); wb.save(out); return out.getvalue()
