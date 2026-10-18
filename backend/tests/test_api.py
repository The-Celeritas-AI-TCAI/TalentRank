import io
from openpyxl import load_workbook
from fastapi.testclient import TestClient
from app.main import app

client=TestClient(app)

def test_health():
    r=client.get('/api/health')
    assert r.status_code==200 and r.json()['status']=='ok'

def test_screening_workflow_generates_excel_and_skips_duplicate():
    role="QA Engineer"
    jd="""Required: Python and SQL. Must have at least 2 years of experience. Nice to have: Docker. Responsibilities include automated testing and API validation."""
    resume="""Nina Example
Email: nina.example@example.com
Phone: +91 98765 43210
Location: Pune, Maharashtra
Professional Experience
QA Engineer — Example Systems
Built Python test automation and SQL validation for API workflows. 3 years of experience.
Skills
Python, SQL, Docker
Projects
Developed a regression test suite for API validation.
Education
Bachelor degree in Computer Science.
"""
    data={"role":role,"job_description":jd}
    files=[
      ("files",("nina.txt",resume.encode(),"text/plain")),
      ("files",("nina-copy.txt",resume.encode(),"text/plain")),
    ]
    response=client.post('/api/screenings',data=data,files=files)
    assert response.status_code==201, response.text
    payload=response.json()
    assert payload["assessed"]==1
    assert len(payload["duplicates"])==1
    screening_id=payload["screening"]["id"]
    candidate=payload["candidates"][0]
    assert candidate["status"]=="scored"
    assert 0 <= candidate["match_score"] <= 100
    assert candidate["profile"]["email"]=="nina.example@example.com"
    assert candidate["profile"]["score_breakdown"]
    assert candidate["profile"]["demonstrated_skills"]
    result=client.get(f'/api/screenings/{screening_id}/results')
    assert result.status_code==200 and len(result.json()["candidates"])==1
    exported=client.get(f'/api/screenings/{screening_id}/export')
    assert exported.status_code==200
    assert exported.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    wb=load_workbook(io.BytesIO(exported.content),data_only=False)
    assert wb.sheetnames==["Candidate Screening","Summary","Rubric"]
    assert wb["Candidate Screening"].auto_filter.ref
    assert wb["Candidate Screening"].freeze_panes=="C5"

def test_extraction_failure_is_reported_not_scored():
    response=client.post('/api/screenings',data={
      "role":"Data Engineer",
      "job_description":"Must have Python and SQL skills. Nice to have: Spark. Responsibilities include data pipelines."
    },files=[("files",("broken.pdf",b"not a real pdf","application/pdf"))])
    assert response.status_code==201,response.text
    candidate=response.json()["candidates"][0]
    assert candidate["status"]=="failed"
    assert candidate["error"]
    assert candidate["match_score"] is None
